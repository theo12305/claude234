import io
import zipfile

import openpyxl
import pytest

from ghe_web import analysis, export, model, plots, properties


@pytest.fixture(scope="module")
def sized():
    inputs = model.default_inputs()
    return inputs, model.size(inputs)


def test_key_results_are_grouped(sized):
    inputs, result = sized
    key = analysis.key_results(result, inputs)
    assert len(key.groups) == len(key.rows)
    assert list(key.by_group()) == ["Temperatures", "Ground energy", "Borehole and hydraulics", "Cost"]
    assert key.values["drift_k"] < 0  # the default load takes more heat out than it puts in
    assert 1000 < key.values["full_load_hours"] < 3000


def test_pressure_drop_is_higher_when_cold(sized):
    inputs, result = sized
    cold = analysis.hydraulics(result, inputs, temperature=0.0)["dp_kpa"]
    warm = analysis.hydraulics(result, inputs, temperature=15.0)["dp_kpa"]
    assert cold > warm


def test_excel_workbook(sized):
    inputs, result = sized
    extra = {"Sensitivity": analysis.sensitivity(inputs)}
    wb = openpyxl.load_workbook(io.BytesIO(export.excel_workbook(result, inputs, "Test", extra)))
    assert wb.sheetnames[:2] == ["Summary", "Inputs"]
    assert {"Monthly temperatures", "Yearly min-max", "Ground energy", "Sensitivity"} <= set(wb.sheetnames)
    years = inputs["load"].values["simulation_period"]
    assert wb["Monthly temperatures"].max_row == 12 * years + 1


def test_zip_bundle(sized):
    inputs, result = sized
    with zipfile.ZipFile(io.BytesIO(export.zip_bundle(result, inputs, "Office"))) as z:
        assert sorted(z.namelist()) == ["Office.json", "Office.pdf", "Office.xlsx"]
        _, loaded = model.from_json(z.read("Office.json").decode())
        assert loaded["ground"].values == inputs["ground"].values


def test_hourly_result_charts_and_export():
    inputs = model.default_inputs()
    inputs["load"].variant = "hourly_building"
    inputs["load"].values = properties.get("load").variant("hourly_building").defaults()
    inputs["load"].values["simulation_period"] = 2
    result = model.size(inputs)
    limits = inputs["limits"].values
    for chart in (plots.hourly_band_chart(result.hourly_temperatures, limits),
                  plots.hourly_last_year_chart(result.hourly_temperatures, limits),
                  plots.heatmap_chart(result.hourly_temperatures)):
        chart.to_json()
    wb = openpyxl.load_workbook(io.BytesIO(export.excel_workbook(result, inputs)), read_only=True)
    assert "Hourly temperatures" in wb.sheetnames and "Input hourly demand" in wb.sheetnames


def test_every_section_has_a_step():
    for section in properties.sections():
        assert section.group, section.key
